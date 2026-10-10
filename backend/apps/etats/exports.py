import csv
import io
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.platypus import Spacer

from apps.editions import pdf as p

LIMITE_PDF = 2000
FORMATS_XLSX = {'montant': '#,##0.00', 'nombre': '#,##0', 'pourcent': '0.0"%"',
                'date': 'DD/MM/YYYY', 'datetime': 'DD/MM/YYYY HH:MM'}


# ------------------------------------------------------------------ Valeurs
def _vide(v):
    return v is None or v == ''


def _texte(type_, v):
    """Valeur pour le CSV : virgule décimale, dates JJ/MM/AAAA."""
    if _vide(v):
        return ''
    if type_ == 'montant':
        return f"{Decimal(v):.2f}".replace('.', ',')
    if type_ == 'pourcent':
        return f"{Decimal(v):.1f}".replace('.', ',')
    if type_ == 'nombre':
        return str(int(v))
    if type_ == 'date':
        return v.strftime('%d/%m/%Y')
    if type_ == 'datetime':
        return timezone.localtime(v).strftime('%d/%m/%Y %H:%M')
    return str(v)


def _ligne_total(etat):
    cols = etat['colonnes']
    ligne = ['' for _ in cols]
    ligne[0] = 'TOTAL'
    for i, c in enumerate(cols):
        if c['cle'] in etat['totaux']:
            ligne[i] = _texte(c['type'], etat['totaux'][c['cle']])
    return ligne


# --------------------------------------------------------------------- CSV
def csv_bytes(etat):
    cols = etat['colonnes']
    sortie = io.StringIO()
    w = csv.writer(sortie, delimiter=';', lineterminator='\r\n')
    w.writerow([c['label'] for c in cols])
    for ligne in etat['lignes']:
        w.writerow([_texte(c['type'], ligne.get(c['cle'])) for c in cols])
    if etat['totaux']:
        w.writerow(_ligne_total(etat))
    return ('\ufeff' + sortie.getvalue()).encode('utf-8')      # BOM : Excel lit correctement l'UTF-8


# -------------------------------------------------------------------- Excel
def _valeur_xlsx(type_, v):
    if _vide(v):
        return None
    if type_ in ('date', 'datetime') and hasattr(v, 'tzinfo') and v.tzinfo is not None:
        return timezone.localtime(v).replace(tzinfo=None)
    return v


def xlsx_bytes(etat):
    cols = etat['colonnes']
    n = len(cols)
    classeur = Workbook()
    ws = classeur.active
    ws.title = 'État'

    ws.append([etat['titre']])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([etat['sous_titre']])
    ws['A2'].font = Font(italic=True, color='555555')
    ws.append([])
    ws.append([c['label'] for c in cols])
    entete = 4
    for i in range(1, n + 1):
        c = ws.cell(row=entete, column=i)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F3A5F')
        c.alignment = Alignment(horizontal='right' if cols[i - 1]['type'] in FORMATS_XLSX and cols[i - 1]['type'] not in ('date', 'datetime') else 'left',
                                vertical='center', wrap_text=True)

    for ligne in etat['lignes']:
        ws.append([_valeur_xlsx(c['type'], ligne.get(c['cle'])) for c in cols])
        for i, c in enumerate(cols, start=1):
            if c['type'] in FORMATS_XLSX:
                ws.cell(row=ws.max_row, column=i).number_format = FORMATS_XLSX[c['type']]

    if etat['totaux']:
        ws.append([_valeur_xlsx(c['type'], etat['totaux'].get(c['cle'])) if i else 'TOTAL'
                   for i, c in enumerate(cols)])
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(row=ws.max_row, column=i)
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='EEF2F7')
            cell.border = Border(top=Side(style='thin', color='1F3A5F'))
            if c['type'] in FORMATS_XLSX:
                cell.number_format = FORMATS_XLSX[c['type']]

    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(11, min(46, int(c['poids'] * 9)))
    ws.freeze_panes = f'A{entete + 1}'
    if etat['lignes']:
        ws.auto_filter.ref = f"A{entete}:{get_column_letter(n)}{entete + len(etat['lignes'])}"

    if etat['synthese']:
        s = classeur.create_sheet('Synthèse')
        s.append([etat['titre'] + ' — synthèse'])
        s['A1'].font = Font(bold=True, size=13)
        s.append([])
        s.append(['', 'Nombre', 'Montant'])
        for i in (1, 2, 3):
            s.cell(row=3, column=i).font = Font(bold=True)
        for l in etat['synthese']:
            s.append([l['label'], l.get('nombre'), l.get('montant')])
            s.cell(row=s.max_row, column=2).number_format = '#,##0'
            s.cell(row=s.max_row, column=3).number_format = '#,##0.00'
        s.column_dimensions['A'].width = 34
        s.column_dimensions['B'].width = 12
        s.column_dimensions['C'].width = 18

    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


# ---------------------------------------------------------------------- PDF
def _montant_pdf(v):
    n = Decimal(v)
    if n == n.to_integral_value():
        return p.nombre(n)
    return f"{n:,.2f}".replace(',', p.NBSP).replace('.', ',')


def _valeur_pdf(type_, v):
    if _vide(v):
        return ''
    if type_ == 'montant':
        return _montant_pdf(v)
    if type_ == 'pourcent':
        return f"{Decimal(v):.1f}".replace('.', ',') + ' %'
    if type_ == 'nombre':
        return p.nombre(v)
    if type_ == 'date':
        return p.fmt_date(v)
    if type_ == 'datetime':
        return p.fmt_dt(v)
    return str(v)


def pdf_bytes(etat, user):
    cols = etat['colonnes']
    largeur = p.LARGEUR_PAYSAGE
    story = [*p.entete(), *p.titre_doc(etat['titre'].upper()), p.P(etat['sous_titre'], 'centre'), Spacer(1, 10)]

    if etat['synthese']:
        lignes = [
            (l['label'], '' if l.get('nombre') is None else p.nombre(l['nombre']),
             '' if l.get('montant') is None else _montant_pdf(l['montant']) + p.NBSP + 'Ar')
            for l in etat['synthese']
        ]
        t = p.tableau(['Synthèse', 'Nombre', 'Montant'], lignes, [50, 15, 30], droite=(1, 2), largeur=largeur * 0.5)
        t.hAlign = 'LEFT'
        story += [t, Spacer(1, 10)]

    lignes = etat['lignes'][:LIMITE_PDF]
    droite = tuple(i for i, c in enumerate(cols) if c['type'] in ('montant', 'nombre', 'pourcent'))
    if lignes:
        story.append(p.tableau(
            [c['label'] for c in cols],
            [[_valeur_pdf(c['type'], l.get(c['cle'])) for c in cols] for l in lignes],
            [c['poids'] for c in cols], droite=droite, largeur=largeur,
        ))
    else:
        story.append(p.P("Aucune donnée pour ces critères.", 'centre'))
    if len(etat['lignes']) > LIMITE_PDF:
        story.append(p.note(
            f"Document limité aux {LIMITE_PDF} premières lignes sur {len(etat['lignes'])} : "
            "utilisez l'export Excel pour disposer de toutes les lignes."
        ))

    if etat['totaux']:
        story += [Spacer(1, 8), p.totaux([
            (f"Total — {c['label']}", _valeur_pdf(c['type'], etat['totaux'][c['cle']]) + (p.NBSP + 'Ar' if c['type'] == 'montant' else ''), True)
            for c in cols if c['cle'] in etat['totaux']
        ])]
    return p.construire(story, etat['titre'], user, paysage=True)